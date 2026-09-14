from __future__ import annotations

import json
from pathlib import Path
from collections.abc import Callable

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None


def build_category_report_sections(category_name: str, file_reviews: list[dict]) -> dict:
    category = category_name.lower()
    relevant = []
    for review in file_reviews:
        file_name = review.get("file", "")
        lower_name = file_name.lower()
        is_test_file = "test" in lower_name or lower_name.endswith(("spec.js", "spec.ts", "spec.tsx"))
        is_ui_file = lower_name.endswith((".html", ".css", ".js", ".jsx", ".ts", ".tsx"))
        is_api_file = lower_name.endswith((".py", ".js", ".jsx", ".ts", ".tsx"))

        if category == "automatic_testing" and is_test_file:
            relevant.append(review)
        elif category == "regression_testing" and (is_test_file or lower_name.endswith(".py")):
            relevant.append(review)
        elif category == "responsiveness_testing" and is_ui_file:
            relevant.append(review)
        elif category == "api_testing" and is_api_file:
            relevant.append(review)
        elif category == "manual_testing" and not lower_name.endswith((".md", ".txt")):
            relevant.append(review)

    if not relevant:
        relevant = file_reviews[:]

    passed = [review for review in relevant if review.get("status") == "passed"]
    needs_work = [review for review in relevant if review.get("status") != "passed"]
    working = [
        f"Static review passed for: {', '.join(review.get('file', '') for review in passed[:8])}."
    ] if passed else ["No relevant file passed the available static checks yet."]
    improvements = []
    for review in needs_work[:8]:
        improvements.append(
            f"{review.get('file', 'Unknown file')}: {' '.join(review.get('improvements', []))}"
        )
    if not improvements:
        improvements.append("No static issue was detected; execute the runtime test and review its result before release.")

    return {
        "working": working,
        "improvements": improvements,
        "guidance": (
            f"This category was prepared, not fully executed. Run the listed command and review the file-level results "
            f"for {len(relevant)} relevant file(s)."
        ),
    }


def build_file_reviews(
    root: Path,
    files: list[str],
    progress_callback: Callable[[int, str], None] | None = None,
) -> list[dict]:
    reviews = []
    total = max(len(files), 1)
    for index, file_name in enumerate(files, start=1):
        if progress_callback:
            progress_callback(round(index / total * 100), f"Static QA checks: {file_name}")
        path = root / file_name
        working = []
        improvements = []
        guidance = []
        extension = path.suffix.lower()

        try:
            text = path.read_text(encoding="utf-8", errors="ignore")
            line_count = len(text.splitlines())
        except OSError as exc:
            reviews.append({
                "file": file_name,
                "status": "skipped",
                "working": [],
                "improvements": [f"File could not be read safely: {exc}"],
                "guidance": "Restore the file or verify archive permissions, then rerun analysis.",
            })
            continue

        if text.strip():
            working.append(f"Readable source/configuration file with {line_count} lines.")
        else:
            improvements.append("The file is empty and currently provides no executable or documented behavior.")

        if extension == ".py":
            try:
                compile(text, str(path), "exec")
                working.append("Python syntax check passed.")
            except SyntaxError as exc:
                improvements.append(f"Python syntax error at line {exc.lineno}: {exc.msg}.")
            guidance.append("Run pytest for this module and add boundary-case tests for public functions.")
        elif extension in {".js", ".jsx", ".ts", ".tsx"}:
            guidance.append("Run the project lint, build, and unit tests; cover loading, error, and empty states.")
        elif extension == ".css":
            guidance.append("Check this stylesheet at mobile, tablet, and desktop breakpoints for overflow and contrast.")
        elif path.name.lower() in {"package.json", "requirements.txt", "pyproject.toml"}:
            guidance.append("Pin reviewed dependency versions and run the project dependency/security audit.")
        else:
            guidance.append("Review this file during the relevant manual and regression QA pass.")

        if "todo" in text.lower() or "fixme" in text.lower():
            improvements.append("TODO/FIXME markers remain and should be converted into tracked work items.")
        if "print(" in text or "console.log(" in text:
            improvements.append("Debug output is present; replace it with structured logging or remove it before release.")

        reviews.append({
            "file": file_name,
            "status": "needs_improvement" if improvements else "passed",
            "working": working or ["File was included in the archive scan."],
            "improvements": improvements or ["No static issue was detected by the current checks."],
            "guidance": " ".join(guidance),
        })
    return reviews


def generate_docx_report(report_path: str, project_name: str, result: dict) -> str:
    path = Path(report_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if Document is None:
        path.write_text(json.dumps(result, indent=2), encoding="utf-8")
        return str(path)

    doc = Document()
    doc.add_heading(f"QA Report - {project_name}", level=1)
    doc.add_paragraph(f"Overall Score: {result.get('overall_score', 0)}")
    doc.add_paragraph(result.get('summary', ''))

    for category_name, category_data in result.get('categories', {}).items():
        doc.add_heading(category_name.replace('_', ' ').title(), level=2)
        doc.add_paragraph(f"Status: {category_data.get('status', 'pending')}")
        doc.add_paragraph(f"Score: {category_data.get('score', 0)}")
        sections = category_data.get("report_sections", {})
        doc.add_paragraph("Verified working points:", style="List Bullet")
        for item in sections.get("working", ["No executed evidence is available yet."]):
            doc.add_paragraph(item)
        doc.add_paragraph("Improvements required:", style="List Bullet")
        for item in sections.get("improvements", category_data.get("details", [])):
            doc.add_paragraph(item)
        doc.add_paragraph("Evidence and execution guidance:", style="List Bullet")
        doc.add_paragraph(sections.get("guidance", "Execute the command below; this status is not a pass result."))
        doc.add_paragraph("How to verify:", style="List Bullet")
        doc.add_paragraph(category_data.get("command", "Perform a focused manual QA review."))

    security = result.get("security", {})
    doc.add_heading("Security Testing Results", level=2)
    doc.add_paragraph(f"Overall security status: {security.get('status', 'unknown')}")
    doc.add_paragraph(f"Security summary: {json.dumps(security.get('summary', {}), default=str)}")
    for scanner_name, scanner_data in security.get("scanners", {}).items():
        doc.add_heading(scanner_name.upper(), level=3)
        doc.add_paragraph(
            f"Status: {scanner_data.get('status', 'unknown')} | "
            f"Findings: {scanner_data.get('summary', {}).get('total', 0)}"
        )
        findings = scanner_data.get("findings", [])
        if findings:
            for finding in findings:
                doc.add_paragraph(json.dumps(finding, default=str), style="List Bullet")
        else:
            doc.add_paragraph("No findings returned by this scanner.", style="List Bullet")
        for error in scanner_data.get("errors", []):
            doc.add_paragraph(f"Scanner note: {error}", style="List Bullet")

    testing = result.get("testing", {})
    doc.add_heading("Testing Results", level=2)
    for test_name, test_data in testing.items():
        doc.add_heading(test_name.replace("_", " ").title(), level=3)
        doc.add_paragraph(
            f"Status: {test_data.get('status', 'unknown')} | "
            f"Cases/findings: {test_data.get('summary', {}).get('total', 0)}"
        )
        for finding in test_data.get("findings", []):
            doc.add_paragraph(json.dumps(finding, default=str), style="List Bullet")
        for error in test_data.get("errors", []):
            doc.add_paragraph(f"Testing note: {error}", style="List Bullet")

    api_testing = result.get("api_testing", {})
    doc.add_heading("API Testing Results", level=2)
    doc.add_paragraph(
        f"Status: {api_testing.get('status', 'unknown')} | "
        f"Endpoints discovered: {api_testing.get('summary', {}).get('total', 0)}"
    )
    for endpoint in api_testing.get("endpoints", []):
        doc.add_paragraph(json.dumps(endpoint, default=str), style="List Bullet")

    runtime = result.get("runtime", {})
    doc.add_heading("Runtime Testing Readiness", level=2)
    doc.add_paragraph(json.dumps(runtime, default=str))
    doc.add_heading("Score Breakdown", level=2)
    doc.add_paragraph(json.dumps(result.get("score_breakdown", {}), default=str))

    file_reviews = result.get("file_reviews", [])
    if file_reviews:
        doc.add_heading("File-by-file QA Review", level=2)
        doc.add_paragraph(
            "Each file below was statically reviewed. A passed result means no issue was found by these checks;"
            " it does not replace runtime or business-flow testing."
        )
        for review in file_reviews:
            doc.add_heading(review.get("file", "Unknown file"), level=3)
            doc.add_paragraph(f"Result: {review.get('status', 'unknown')}")
            doc.add_paragraph("Working points:", style="List Bullet")
            for item in review.get("working", []):
                doc.add_paragraph(item)
            doc.add_paragraph("Improvements:", style="List Bullet")
            for item in review.get("improvements", []):
                doc.add_paragraph(item)
            doc.add_paragraph(f"Guidance: {review.get('guidance', '')}")

    doc.save(path)
    return str(path)


def generate_json_report(report_path: str, result: dict) -> str:
    path = Path(report_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    return str(path)


def generate_html_report(report_path: str, project_name: str, result: dict) -> str:
    path = Path(report_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    summary = result.get("summary", "")
    score = result.get("overall_score", 0)
    security = result.get("security", {}).get("summary", {})
    body = (
        f"<h1>QA Report - {project_name}</h1>"
        f"<p><strong>Overall score:</strong> {score}</p>"
        f"<p>{summary}</p>"
        f"<h2>Security</h2><p>{json.dumps(security, default=str)}</p>"
        f"<h2>Evidence</h2><pre>{json.dumps(result, indent=2, default=str)}</pre>"
    )
    path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><title>QA Report</title></head><body>{body}</body></html>", encoding="utf-8")
    return str(path)
